import logging
from pathlib import Path
from typing import Any
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from domain.errors import DomainError
from domain.interfaces import IReportExporter

logger = logging.getLogger(__name__)

class ExportError(DomainError):
    pass

class DocxExporter(IReportExporter):
    def export(self, report_data: dict[str, Any], output_path: str) -> str:
        logger.debug(f"{self.__class__.__name__}.export: {output_path}")
        doc = Document()
        self._add_title(doc, report_data.get("title", "Аналитический отчёт"))
        self._add_text(doc, report_data.get("insights", ""))
        self._add_table(doc, report_data.get("statistics", {}))
        self._add_images(doc, report_data.get("chart_images", []))
        doc.save(output_path)
        return output_path

    def _add_title(self, doc: Document, title: str):
        doc.add_heading(title, level=0).alignment = WD_ALIGN_PARAGRAPH.CENTER

    def _add_text(self, doc: Document, text: str):
        doc.add_paragraph(text).style.font.size = Pt(11)

    def _add_table(self, doc: Document, stats: dict[str, dict]):
        if not stats:
            return
        table = doc.add_table(rows=1, cols=5)
        table.style = "Light Shading Accent 1"
        for i, h in enumerate(["Колонка", "Пропуски", "Уникальные", "Среднее", "Макс"]):
            table.rows[0].cells[i].text = h
        for col, data in stats.items():
            row = table.add_row()
            row.cells[0].text = col
            for i, k in enumerate(["missing", "unique", "mean", "max"], 1):
                row.cells[i].text = str(data.get(k, "-"))

    def _add_images(self, doc: Document, images: list):
        for item in images:
            path = item.get("image_path") if isinstance(item, dict) else item
            title = item.get("title", "") if isinstance(item, dict) else ""
            if path and Path(path).exists():
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(path, width=Inches(5.5))
                if title:
                    doc.add_paragraph(title).alignment = WD_ALIGN_PARAGRAPH.CENTER